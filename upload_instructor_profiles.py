# upload_instructor_profiles.py
# 실행: python upload_instructor_profiles.py
# 위치: C:\pnl-dashboard\
# 동작:
#   1. Supabase Storage(teachers 버킷)에 사진 업로드
#   2. instructor_profiles 테이블 전체 재업로드 (360명)

import os
from openpyxl import load_workbook
from datetime import datetime
from supabase import create_client
from dotenv import load_dotenv

load_dotenv(dotenv_path=".env.local")

SUPABASE_URL = os.getenv("NEXT_PUBLIC_SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
supabase     = create_client(SUPABASE_URL, SUPABASE_KEY)

EXCEL_PATH = r"C:\Users\whydu\Desktop\data-upload\강사프로필.xlsx"
PHOTO_DIR  = r"C:\Users\whydu\Desktop\data-upload\teachers"
BUCKET     = "teachers"

def parse_excel():
    wb = load_workbook(EXCEL_PATH, read_only=True)
    ws = wb['강사정보']
    rows = list(ws.iter_rows(values_only=True))
    data = [r for r in rows[1:] if any(v is not None for v in r)]

    records = []
    for r in data:
        file_raw = str(r[0]).strip() if r[0] else None
        name     = str(r[1]).strip() if r[1] else None
        if not name:
            continue

        file_key = file_raw.replace('.png','').strip() if file_raw else None
        has_photo = file_raw and str(file_raw).endswith('.png')
        photo_url = (
            f"{SUPABASE_URL}/storage/v1/object/public/{BUCKET}/{file_key}.png"
            if has_photo else None
        )

        birth_year = None
        if r[3]:
            try: birth_year = int(r[3])
            except: pass

        def fmt_date(v):
            if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
            if v: return str(v)[:10]
            return None

        records.append({
            'file_key':       file_key,
            'name':           name,
            'name_kr':        name,
            'gender':         str(r[2]).strip() if r[2] else None,
            'birth_year':     birth_year,
            'phone':          str(r[4]).strip() if r[4] else None,
            'subject':        str(r[5]).strip() if r[5] else None,
            'subjects':       str(r[5]).strip() if r[5] else None,
            'academy':        str(r[6]).strip() if r[6] else None,
            'education':      str(r[7]).strip() if r[7] else None,
            'contract_type':  str(r[8]).strip() if r[8] else None,
            'contract_start': fmt_date(r[9]),
            'contract_end':   fmt_date(r[10]),
            'memo':           str(r[11]).strip() if r[11] else None,
            'status':         str(r[12]).strip() if r[12] else '출강',
            'photo_url':      photo_url,
        })
    return records

def upload_photos(records):
    if not os.path.exists(PHOTO_DIR):
        print(f"⚠ 사진 폴더 없음: {PHOTO_DIR} (사진 없이 계속)")
        return

    files = {f.replace('.png','').lower(): f
             for f in os.listdir(PHOTO_DIR) if f.lower().endswith('.png')}
    print(f"[1/2] 사진 업로드 중 ({len(files)}개)...")

    for i, (key, fname) in enumerate(sorted(files.items())):
        fpath = os.path.join(PHOTO_DIR, fname)
        with open(fpath, 'rb') as f:
            content = f.read()
        try:
            supabase.storage.from_(BUCKET).upload(
                path=fname,
                file=content,
                file_options={"content-type": "image/png", "upsert": "true"}
            )
        except Exception as e:
            if "already exists" not in str(e).lower():
                print(f"\n  ⚠ {fname}: {e}")
        print(f"  {i+1}/{len(files)}", end="\r")
    print(f"\n  업로드 완료\n")

def insert_profiles(records):
    print(f"[2/2] instructor_profiles {len(records)}명 업로드 중...")
    # 기존 전체 삭제
    supabase.table("instructor_profiles").delete().neq("id", 0).execute()

    BATCH = 50
    for i in range(0, len(records), BATCH):
        supabase.table("instructor_profiles").insert(records[i:i+BATCH]).execute()
        print(f"  {min(i+BATCH, len(records))} / {len(records)}", end="\r")
    print(f"\n  INSERT 완료\n")

def main():
    print("=== 강사 프로필 업로드 시작 ===\n")
    records = parse_excel()
    print(f"파싱 완료: {len(records)}명 "
          f"(사진 {sum(1 for r in records if r['photo_url'])}명 / "
          f"사진없음 {sum(1 for r in records if not r['photo_url'])}명)\n")

    upload_photos(records)
    insert_profiles(records)
    print("✅ 완료!")

if __name__ == "__main__":
    main()
