# upload_admission.py
# 실행: python upload_admission.py
# 위치: C:\pnl-dashboard\
# 동작: admission_results 테이블 전체 삭제 후 2024~2026 재업로드
# 남-의대관 → 남학생 으로 자동 변환

import os
from openpyxl import load_workbook
from supabase import create_client
from dotenv import load_dotenv

load_dotenv(dotenv_path=".env.local")

SUPABASE_URL = os.getenv("NEXT_PUBLIC_SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise ValueError("환경변수 NEXT_PUBLIC_SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY 를 확인하세요.")

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

EXCEL_PATH = r"C:\Users\whydu\Desktop\data-upload\2024-2026학년도_입결현황_최종.xlsx"
SHEETS     = ["2026학년도", "2025학년도", "2024학년도"]

ACADEMY_NORMALIZE = {
    "남-의대관": "남학생",
}

def parse_sheet(ws) -> list:
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    records = []
    for row in rows[1:]:
        if not any(v is not None for v in row):
            continue
        r = dict(zip(header, row))
        academy = (r.get("학원") or "").strip()
        academy = ACADEMY_NORMALIZE.get(academy, academy)
        records.append({
            "year":           int(r["학년도"]) if r.get("학년도") else None,
            "academy":        academy,
            "admission_type": (r.get("수시/정시") or "").strip() or None,
            "university":     (r.get("대학") or "").strip() or None,
            "univ_group":     (r.get("대학그룹") or "").strip() or None,
            "univ_group2":    (r.get("대학그룹2") or "").strip() or None,
        })
    return records

def main():
    wb = load_workbook(EXCEL_PATH, read_only=True)

    all_records = []
    for sname in SHEETS:
        ws = wb[sname]
        records = parse_sheet(ws)
        print(f"  {sname}: {len(records)}건 파싱 완료")
        all_records.extend(records)

    print(f"\n총 {len(all_records)}건 파싱 완료")

    print("\n[1/2] 기존 admission_results 전체 삭제 중...")
    supabase.table("admission_results").delete().neq("year", 0).execute()
    print("  삭제 완료")

    print(f"[2/2] {len(all_records)}건 업로드 중...")
    BATCH = 500
    for i in range(0, len(all_records), BATCH):
        batch = all_records[i:i+BATCH]
        supabase.table("admission_results").insert(batch).execute()
        print(f"  {min(i+BATCH, len(all_records))} / {len(all_records)} 완료", end="\r")

    print(f"\n✅ 업로드 완료: 총 {len(all_records)}건")

if __name__ == "__main__":
    main()
